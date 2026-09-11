from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
import os, datetime

PLOT_DIR = "plots"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(0x1E, 0x27, 0x61)
ICE  = RGBColor(0xCA, 0xDC, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DKGRAY = RGBColor(0x33, 0x33, 0x33)
LTGRAY = RGBColor(0xF2, 0xF2, 0xF2)
FAIL_RED = RGBColor(0xB8, 0x50, 0x42)
PASS_GREEN = RGBColor(0x2C, 0x5F, 0x2D)
COND_AMBER = RGBColor(0xD4, 0x8B, 0x0B)
WORSE_BG = RGBColor(0xFC, 0xE4, 0xE4)
SAME_BG = RGBColor(0xE8, 0xF5, 0xE9)
BLACK = RGBColor(0x00, 0x00, 0x00)

def set_cell_bg(cell, color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    solidFill = tcPr.makeelement(qn("a:solidFill"), {})
    srgbClr = solidFill.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % (color[0], color[1], color[2])})
    solidFill.append(srgbClr)
    tcPr.append(solidFill)

def style_cell(cell, text, font_size=12, bold=False, color=DKGRAY, align=PP_ALIGN.CENTER, bg=None):
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if bg:
        set_cell_bg(cell, bg)

# ================================================================
# SLIDE 1: Title
# ================================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = NAVY
tx = sl.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(3.5))
tf = tx.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "National Lab of Rockies\nGFM Inverter Model\nMQT Evaluation \u2014 Anti-Windup Comparison"
p.font.size = Pt(38); p.font.bold = True; p.font.color.rgb = WHITE
p2 = tf.add_paragraph()
p2.text = "ERCOT DWG Procedure Manual Rev. 24\nSOW Task 2 \u2014 PMView 2.4\nVoltage Regulation ON | Pset = 0.4 pu"
p2.font.size = Pt(16); p2.font.color.rgb = ICE; p2.space_before = Pt(24)
p3 = tf.add_paragraph()
p3.text = datetime.datetime.now().strftime("%m/%d/%Y %I:%M %p")
p3.font.size = Pt(14); p3.font.color.rgb = ICE; p3.space_before = Pt(12)

# ================================================================
# SLIDE 2: Combined Summary Table
# ================================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
# Title bar
tb = sl.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(1.0))
tb.fill.solid(); tb.fill.fore_color.rgb = NAVY; tb.line.fill.background()
ttx = tb.text_frame; ttx.word_wrap = True; ttx.margin_left = Inches(0.6)
p = ttx.paragraphs[0]
p.text = "MQT Results Summary \u2014 No Anti-Windup vs. Anti-Windup ON"
p.font.size = Pt(26); p.font.bold = True; p.font.color.rgb = WHITE

# Table: 11 rows x 7 cols
rows, cols = 11, 7
tbl = sl.shapes.add_table(rows, cols, Inches(0.4), Inches(1.2), Inches(12.5), Inches(5.8)).table

# Column widths
widths = [0.6, 2.8, 0.8, 2.8, 1.8, 1.8, 1.5]
for i, w in enumerate(widths):
    tbl.columns[i].width = Inches(w)

# Header row
headers = ["Test", "Name", "Section", "Configuration", "No Anti-Windup", "Anti-Windup ON", "Change"]
for i, h in enumerate(headers):
    style_cell(tbl.cell(0, i), h, font_size=11, bold=True, color=WHITE, bg=(0x1E, 0x27, 0x61))

# Data
test_data = [
    ("1", "Flat Start", "3.1.5.2",
     "Vsched=1.03 | V-Reg ON", "PASS", "PASS", "Same"),
    ("2", "LVRT ERCOT Legacy", "3.1.5.4",
     "Vsched=1.03 | V-Reg ON", "PASS", "FAIL", "Worse"),
    ("3", "HVRT ERCOT Legacy", "3.1.5.5",
     "Vsched=1.03 | V-Reg ON", "COND. PASS", "FAIL", "Worse"),
    ("4", "Small V Disturbance (Step Down)", "3.1.5.3",
     "Vsched=1.01 | V-Reg ON", "PASS", "PASS", "Same"),
    ("5", "Small V Disturbance (Step Up)", "3.1.5.3",
     "Vsched=1.01 | V-Reg ON", "PASS", "PASS", "Same"),
    ("6", "LVRT Dips IEEE 2800", "3.1.5.4",
     "Vsched=1.03 | V-Reg ON", "PASS", "COND. PASS", "Worse"),
    ("7", "HVRT Preferred IEEE 2800", "3.1.5.5",
     "Vsched=1.03 | V-Reg ON", "COND. PASS", "FAIL", "Worse"),
    ("8", "Phase Angle Jump (Down)", "3.1.5.9",
     "Vsched=1.01 | V-Reg ON", "FAIL", "FAIL", "Same"),
    ("9", "Phase Angle Jump (Up)", "3.1.5.9",
     "Vsched=1.01 | V-Reg ON", "FAIL", "FAIL", "Same"),
]

def result_color(txt):
    if "FAIL" == txt: return FAIL_RED
    if "COND" in txt: return COND_AMBER
    if "PASS" == txt: return PASS_GREEN
    return DKGRAY

def change_bg(txt):
    if txt == "Worse": return WORSE_BG
    if txt == "Same": return SAME_BG
    return None

for r, (tnum, name, sec, config, no_aw, aw, change) in enumerate(test_data, 1):
    row_bg = LTGRAY if r % 2 == 0 else None
    style_cell(tbl.cell(r, 0), tnum, font_size=11, bold=True, bg=row_bg)
    style_cell(tbl.cell(r, 1), name, font_size=10, align=PP_ALIGN.LEFT, bg=row_bg)
    style_cell(tbl.cell(r, 2), sec, font_size=10, bg=row_bg)
    style_cell(tbl.cell(r, 3), config, font_size=9, align=PP_ALIGN.LEFT, bg=row_bg)
    style_cell(tbl.cell(r, 4), no_aw, font_size=11, bold=True, color=result_color(no_aw), bg=row_bg)
    style_cell(tbl.cell(r, 5), aw, font_size=11, bold=True, color=result_color(aw), bg=row_bg)

    cbg = change_bg(change)
    chg_color = FAIL_RED if change == "Worse" else PASS_GREEN
    style_cell(tbl.cell(r, 6), change, font_size=11, bold=True, color=chg_color, bg=cbg)

# Summary row
style_cell(tbl.cell(10, 0), "", bg=(0x1E, 0x27, 0x61))
style_cell(tbl.cell(10, 1), "", bg=(0x1E, 0x27, 0x61))
style_cell(tbl.cell(10, 2), "", bg=(0x1E, 0x27, 0x61))
style_cell(tbl.cell(10, 3), "Overall:", font_size=12, bold=True, color=WHITE,
           align=PP_ALIGN.RIGHT, bg=(0x1E, 0x27, 0x61))
style_cell(tbl.cell(10, 4), "5P / 2C / 2F", font_size=12, bold=True, color=WHITE,
           bg=(0x1E, 0x27, 0x61))
style_cell(tbl.cell(10, 5), "3P / 1C / 5F", font_size=12, bold=True, color=WHITE,
           bg=(0x1E, 0x27, 0x61))
style_cell(tbl.cell(10, 6), "4 Worse / 5 Same", font_size=10, bold=True,
           color=FAIL_RED, bg=WORSE_BG)

# ================================================================
# SLIDE 3: Key Findings
# ================================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
tb = sl.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(1.0))
tb.fill.solid(); tb.fill.fore_color.rgb = NAVY; tb.line.fill.background()
ttx = tb.text_frame; ttx.word_wrap = True; ttx.margin_left = Inches(0.6)
p = ttx.paragraphs[0]
p.text = "Key Findings \u2014 Anti-Windup Current Limiting Impact"
p.font.size = Pt(26); p.font.bold = True; p.font.color.rgb = WHITE

cx = sl.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12.0), Inches(5.8))
cf = cx.text_frame; cf.word_wrap = True

findings = [
    ("Tests Unchanged (1, 4, 5, 8, 9):",
     "Tests where the current limiter does not activate (flat start, small-signal voltage "
     "steps) show identical behavior to the no-anti-windup baseline. Phase angle jump tests "
     "(8, 9) remain FAIL in both cases \u2014 the current limiter does not address the "
     "fundamental power-angle instability."),
    ("Tests Degraded \u2014 HVRT (3, 7):",
     "Both HVRT tests show sustained P/Q oscillations during the entire ride-through period "
     "when anti-windup is enabled. The current limiter cycles on/off repeatedly, creating a "
     "limit-cycle oscillation. This is worse than the no-anti-windup case, which had brief "
     "transient spikes but no sustained oscillations. Tests 3 and 7 move from CONDITIONAL "
     "PASS to FAIL."),
    ("Tests Degraded \u2014 LVRT (2, 6):",
     "LVRT ERCOT Legacy (Test 2) develops growing oscillations during the 0.9 pu sustained "
     "segment and fails to recover properly (Ppoi reverses to -0.46 pu), moving from PASS "
     "to FAIL. LVRT Dips IEEE 2800 (Test 6) rides through all dips but with larger transient "
     "excursions and a brief frequency dip to 53 Hz, moving from PASS to CONDITIONAL PASS."),
    ("Root Cause:",
     "The anti-windup current limiter appears to introduce a limit-cycle instability when "
     "activated during voltage events. The limiter engages, reduces current, which changes "
     "the voltage/power balance, causing the limiter to disengage, which triggers another "
     "excursion, re-engaging the limiter. This cycling pattern is most severe during "
     "sustained over-voltage (HVRT) events where the limiter is continuously near its "
     "threshold."),
]

first = True
for heading, body in findings:
    if first:
        ph = cf.paragraphs[0]; first = False
    else:
        ph = cf.add_paragraph()
    ph.text = heading
    ph.font.size = Pt(14); ph.font.bold = True; ph.font.color.rgb = NAVY
    ph.space_before = Pt(16)

    pb = cf.add_paragraph()
    pb.text = body
    pb.font.size = Pt(12); pb.font.color.rgb = DKGRAY
    pb.space_before = Pt(4)

# ================================================================
# Save
# ================================================================
ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
fname = "GFM_MQT_AW_Summary_Comparison_%s.pptx" % ts
prs.save(fname)
print("Saved: %s (%d slides)" % (fname, len(prs.slides)))
