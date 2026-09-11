"""
Create a standalone PPTX with just the comparison table,
using the original deck as template for theme preservation.
Delete all original slides, keep only the new one.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
import datetime

SRC = "NLR_GFM_MQT_Evaluation_20260406_220436.pptx"
prs = Presentation(SRC)

# Use slide 2's layout
layout = prs.slides[1].slide_layout

# Delete all existing slides
while len(prs.slides) > 0:
    rId = prs.slides._sldIdLst[0].get(qn("r:id"))
    prs.part.drop_rel(rId)
    prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])

# Colors
HDR_BG = "1E3A5F"
ALT_BG = "E8F0FE"
SUM_BG = "D0D8E8"
PASS_COL = RGBColor(0x22, 0x8B, 0x22)
COND_COL = RGBColor(0xCC, 0x8C, 0x00)
FAIL_COL = RGBColor(0xCC, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SUM_TEXT = RGBColor(0x1E, 0x3A, 0x5F)
FONT_NAME = "Tw Cen MT"
HDR_SIZE = Emu(139700)
DATA_SIZE = Emu(127000)
TITLE_SIZE = Emu(228600)


def set_cell_fill(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("a:solidFill")):
        tcPr.remove(old)
    for old in tcPr.findall(qn("a:noFill")):
        tcPr.remove(old)
    solidFill = tcPr.makeelement(qn("a:solidFill"), {})
    srgbClr = solidFill.makeelement(qn("a:srgbClr"), {"val": hex_color})
    solidFill.append(srgbClr)
    tcPr.append(solidFill)


def set_cell_no_fill(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("a:solidFill")):
        tcPr.remove(old)


def style_cell(cell, text, size=None, bold=None, color=None, align=PP_ALIGN.CENTER,
               bg_hex=None):
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = FONT_NAME
    if size:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = color
    p.alignment = align
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if bg_hex:
        set_cell_fill(cell, bg_hex)
    else:
        set_cell_no_fill(cell)


def result_color(txt):
    if txt == "PASS":
        return PASS_COL
    elif "CONDITIONAL" in txt or "COND" in txt:
        return COND_COL
    elif txt == "FAIL":
        return FAIL_COL
    return None


# ── Slide 1: Title (matching original slide 1 style) ──
sl1 = prs.slides.add_slide(layout)
for shape in sl1.shapes:
    if shape.has_text_frame and shape.name.startswith("Title"):
        shape.text_frame.paragraphs[0].text = ""
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = "MQT Results Summary \u2014 No Anti-Windup vs. Anti-Windup ON"
        run.font.name = FONT_NAME
        run.font.size = TITLE_SIZE
    if shape.has_text_frame and shape.name.startswith("Content"):
        sp = shape._element
        sp.getparent().remove(sp)

# Table: 11 rows x 6 cols
rows, cols = 11, 6
tbl_shape = sl1.shapes.add_table(rows, cols,
    Inches(0.35), Inches(1.30), Inches(9.40), Inches(3.80))
table = tbl_shape.table

col_widths = [0.5, 2.5, 0.8, 2.6, 1.5, 1.5]
for i, w in enumerate(col_widths):
    table.columns[i].width = Inches(w)
for ri in range(rows):
    table.rows[ri].height = Emu(314325)

# Header
headers = ["Test", "Name", "Section", "Configuration", "No Anti-Windup", "Anti-Windup ON"]
for ci, h in enumerate(headers):
    style_cell(table.cell(0, ci), h, size=HDR_SIZE, bold=True, color=WHITE, bg_hex=HDR_BG)

# Data
test_data = [
    ("1", "Flat Start", "3.1.5.2", "Vsched = 1.03 | V-Reg ON", "PASS", "PASS"),
    ("2", "LVRT ERCOT Legacy", "3.1.5.4", "Vsched = 1.03 | V-Reg ON", "PASS", "FAIL"),
    ("3", "HVRT ERCOT Legacy", "3.1.5.5", "Vsched = 1.03 | V-Reg ON", "CONDITIONAL PASS", "FAIL"),
    ("4", "Small V Disturbance (Step Down)", "3.1.5.3", "Vsched = 1.01 | V-Reg ON", "PASS", "PASS"),
    ("5", "Small V Disturbance (Step Up)", "3.1.5.3", "Vsched = 1.01 | V-Reg ON", "PASS", "PASS"),
    ("6", "LVRT Dips IEEE 2800", "3.1.5.4", "Vsched = 1.03 | V-Reg ON", "PASS", "CONDITIONAL PASS"),
    ("7", "HVRT Preferred IEEE 2800", "3.1.5.5", "Vsched = 1.03 | V-Reg ON", "CONDITIONAL PASS", "FAIL"),
    ("8", "Phase Angle Jump (Down)", "3.1.5.9", "Vsched = 1.01 | V-Reg ON", "FAIL", "FAIL"),
    ("9", "Phase Angle Jump (Up)", "3.1.5.9", "Vsched = 1.01 | V-Reg ON", "FAIL", "FAIL"),
]

for r, (tnum, name, sec, config, no_aw, aw) in enumerate(test_data, 1):
    row_bg = ALT_BG if r % 2 == 1 else None
    style_cell(table.cell(r, 0), tnum, size=DATA_SIZE, bg_hex=row_bg)
    style_cell(table.cell(r, 1), name, size=DATA_SIZE, align=PP_ALIGN.LEFT, bg_hex=row_bg)
    style_cell(table.cell(r, 2), sec, size=DATA_SIZE, bg_hex=row_bg)
    style_cell(table.cell(r, 3), config, size=DATA_SIZE, bg_hex=row_bg)
    style_cell(table.cell(r, 4), no_aw, size=DATA_SIZE, bold=True,
               color=result_color(no_aw), bg_hex=row_bg)
    style_cell(table.cell(r, 5), aw, size=DATA_SIZE, bold=True,
               color=result_color(aw), bg_hex=row_bg)

# Summary row
for ci in range(4):
    style_cell(table.cell(10, ci), "", bg_hex=SUM_BG)
style_cell(table.cell(10, 3), "Overall:", size=DATA_SIZE, bold=True, bg_hex=SUM_BG)
style_cell(table.cell(10, 4), "5P / 2C / 2F", size=DATA_SIZE, bold=True,
           color=SUM_TEXT, bg_hex=SUM_BG)
style_cell(table.cell(10, 5), "3P / 1C / 5F", size=DATA_SIZE, bold=True,
           color=SUM_TEXT, bg_hex=SUM_BG)

# Page number
txBox = sl1.shapes.add_textbox(Inches(9.25), Inches(5.147), Inches(0.681), Inches(0.324))
p = txBox.text_frame.paragraphs[0]
run = p.add_run()
run.text = "1"
run.font.name = FONT_NAME
run.font.size = DATA_SIZE
p.alignment = PP_ALIGN.RIGHT

# Save
ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
fname = "GFM_MQT_AW_Comparison_Table_%s.pptx" % ts
prs.save(fname)
print("Saved: %s (%d slides)" % (fname, len(prs.slides)))
