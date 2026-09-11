"""
Regenerate the MQT summary comparison table using the exact formatting
from slide 2 of NLR_GFM_MQT_Evaluation_20260406_220436.pptx.

Slide: 10.0 x 5.625 in, Layout: "Title and Content"
Font: Tw Cen MT throughout
Header: 11pt bold white on #1E3A5F
Data: 10pt, alternating #E8F0FE / no fill
Result colors: PASS=#228B22, COND PASS=#CC8C00, FAIL=#CC0000
Summary row: bg #D0D8E8
Table: left=0.35, top=1.30, w=9.40, h=3.80
Cols: 0.5, 3.0, 1.0, 3.2, 1.7
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
import os, datetime, copy

SRC = "NLR_GFM_MQT_Evaluation_20260406_220436.pptx"

# Use original as template to preserve theme/layouts
prs = Presentation(SRC)

# Use the same layout as slide 2 from the original
layout = prs.slides[1].slide_layout

# Colors - exact from original
HDR_BG = "1E3A5F"
ALT_BG = "E8F0FE"
SUM_BG = "D0D8E8"
PASS_COL = RGBColor(0x22, 0x8B, 0x22)
COND_COL = RGBColor(0xCC, 0x8C, 0x00)
FAIL_COL = RGBColor(0xCC, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HDR_BG_RGB = RGBColor(0x1E, 0x3A, 0x5F)
SUM_TEXT = RGBColor(0x1E, 0x3A, 0x5F)

FONT_NAME = "Tw Cen MT"
HDR_SIZE = Emu(139700)   # 11pt
DATA_SIZE = Emu(127000)  # 10pt
TITLE_SIZE = Emu(228600) # 18pt


def set_cell_fill(cell, hex_color):
    """Set solid fill on a cell using hex color string."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    # Remove existing fills
    for old in tcPr.findall(qn("a:solidFill")):
        tcPr.remove(old)
    for old in tcPr.findall(qn("a:noFill")):
        tcPr.remove(old)
    solidFill = tcPr.makeelement(qn("a:solidFill"), {})
    srgbClr = solidFill.makeelement(qn("a:srgbClr"), {"val": hex_color})
    solidFill.append(srgbClr)
    tcPr.append(solidFill)


def set_cell_no_fill(cell):
    """Remove fill from cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("a:solidFill")):
        tcPr.remove(old)


def style_cell(cell, text, size=None, bold=None, color=None, align=PP_ALIGN.CENTER,
               bg_hex=None, font_name=FONT_NAME):
    """Style a cell to match original formatting."""
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    if size:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = color
    p.alignment = align
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    if bg_hex:
        set_cell_fill(cell, bg_hex)
    else:
        set_cell_no_fill(cell)


def result_color(txt):
    if txt == "PASS":
        return PASS_COL
    elif "CONDITIONAL" in txt or "COND" in txt:
        return COND_COL
    elif txt == "FAIL":
        return FAIL_COL
    return None


# ================================================================
# Add new slide at the end using the same layout
# ================================================================
new_slide = prs.slides.add_slide(layout)

# Set title
for shape in new_slide.shapes:
    if shape.has_text_frame and shape.name.startswith("Title"):
        shape.text_frame.paragraphs[0].text = ""
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = "MQT Results Summary \u2014 No Anti-Windup vs. Anti-Windup ON"
        run.font.name = FONT_NAME
        run.font.size = TITLE_SIZE
        break

# Create table: 11 rows x 6 cols
# Original was 5 cols; we add a 6th for "Anti-Windup ON" result
# Adjust column widths to fit the 6 cols in 9.4 inches
rows, cols = 11, 6
left = Inches(0.35)
top = Inches(1.30)
width = Inches(9.40)
height = Inches(3.80)

tbl_shape = new_slide.shapes.add_table(rows, cols, left, top, width, height)
table = tbl_shape.table

# Column widths: Test(0.5) Name(2.5) Section(0.8) Config(2.6) NoAW(1.5) AW(1.5)
col_widths = [0.5, 2.5, 0.8, 2.6, 1.5, 1.5]
for i, w in enumerate(col_widths):
    table.columns[i].width = Inches(w)

# Row heights
for ri in range(rows):
    table.rows[ri].height = Emu(314325)  # ~0.345 in

# Header row
headers = ["Test", "Name", "Section", "Configuration", "No Anti-Windup", "Anti-Windup ON"]
for ci, h in enumerate(headers):
    style_cell(table.cell(0, ci), h, size=HDR_SIZE, bold=True, color=WHITE, bg_hex=HDR_BG)

# Data
test_data = [
    ("1", "Flat Start", "3.1.5.2",
     "Vsched = 1.03 | V-Reg ON", "PASS", "PASS"),
    ("2", "LVRT ERCOT Legacy", "3.1.5.4",
     "Vsched = 1.03 | V-Reg ON", "PASS", "FAIL"),
    ("3", "HVRT ERCOT Legacy", "3.1.5.5",
     "Vsched = 1.03 | V-Reg ON", "CONDITIONAL PASS", "FAIL"),
    ("4", "Small V Disturbance (Step Down)", "3.1.5.3",
     "Vsched = 1.01 | V-Reg ON", "PASS", "PASS"),
    ("5", "Small V Disturbance (Step Up)", "3.1.5.3",
     "Vsched = 1.01 | V-Reg ON", "PASS", "PASS"),
    ("6", "LVRT Dips IEEE 2800", "3.1.5.4",
     "Vsched = 1.03 | V-Reg ON", "PASS", "CONDITIONAL PASS"),
    ("7", "HVRT Preferred IEEE 2800", "3.1.5.5",
     "Vsched = 1.03 | V-Reg ON", "CONDITIONAL PASS", "FAIL"),
    ("8", "Phase Angle Jump (Down)", "3.1.5.9",
     "Vsched = 1.01 | V-Reg ON", "FAIL", "FAIL"),
    ("9", "Phase Angle Jump (Up)", "3.1.5.9",
     "Vsched = 1.01 | V-Reg ON", "FAIL", "FAIL"),
]

for r, (tnum, name, sec, config, no_aw, aw) in enumerate(test_data, 1):
    # Alternating row background: odd rows (1,3,5,7,9) get ALT_BG
    row_bg = ALT_BG if r % 2 == 1 else None

    style_cell(table.cell(r, 0), tnum, size=DATA_SIZE, bg_hex=row_bg)
    style_cell(table.cell(r, 1), name, size=DATA_SIZE, align=PP_ALIGN.LEFT, bg_hex=row_bg)
    style_cell(table.cell(r, 2), sec, size=DATA_SIZE, bg_hex=row_bg)
    style_cell(table.cell(r, 3), config, size=DATA_SIZE, bg_hex=row_bg)

    # No Anti-Windup result
    style_cell(table.cell(r, 4), no_aw, size=DATA_SIZE, bold=True,
               color=result_color(no_aw), bg_hex=row_bg)

    # Anti-Windup ON result
    style_cell(table.cell(r, 5), aw, size=DATA_SIZE, bold=True,
               color=result_color(aw), bg_hex=row_bg)

# Summary row (row 10)
for ci in range(4):
    style_cell(table.cell(10, ci), "", bg_hex=SUM_BG)
style_cell(table.cell(10, 3), "Overall:", size=DATA_SIZE, bold=True, bg_hex=SUM_BG)
style_cell(table.cell(10, 4), "5P / 2C / 2F", size=DATA_SIZE, bold=True,
           color=SUM_TEXT, bg_hex=SUM_BG)
style_cell(table.cell(10, 5), "3P / 1C / 5F", size=DATA_SIZE, bold=True,
           color=SUM_TEXT, bg_hex=SUM_BG)

# Page number textbox (match original position)
txBox = new_slide.shapes.add_textbox(Inches(9.25), Inches(5.147), Inches(0.681), Inches(0.324))
tf = txBox.text_frame
p = tf.paragraphs[0]
run = p.add_run()
run.text = str(len(prs.slides))
run.font.name = FONT_NAME
run.font.size = DATA_SIZE
p.alignment = PP_ALIGN.RIGHT

# Remove the "Click to add text" content placeholder if present
for shape in new_slide.shapes:
    if shape.has_text_frame and shape.name.startswith("Content"):
        sp = shape._element
        sp.getparent().remove(sp)

# Save
ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
fname = "GFM_MQT_AW_vs_NoAW_Comparison_%s.pptx" % ts
prs.save(fname)
print("Saved: %s" % fname)
print("New comparison table is on slide %d" % len(prs.slides))
