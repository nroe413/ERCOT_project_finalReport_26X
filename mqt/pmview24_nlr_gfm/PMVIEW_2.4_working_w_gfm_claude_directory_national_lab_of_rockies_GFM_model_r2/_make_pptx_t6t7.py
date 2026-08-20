from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os, datetime

PLOT_DIR = "plots"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(0x1E, 0x27, 0x61)
ICE  = RGBColor(0xCA, 0xDC, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DKGRAY = RGBColor(0x33, 0x33, 0x33)
FAIL_RED = RGBColor(0xB8, 0x50, 0x42)
PASS_GREEN = RGBColor(0x2C, 0x5F, 0x2D)
COND_AMBER = RGBColor(0xD4, 0x8B, 0x0B)

def add_title_slide(title, subtitle):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = NAVY
    tx = sl.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(2.5))
    tf = tx.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title
    p.font.size = Pt(40); p.font.bold = True; p.font.color.rgb = WHITE
    p2 = tf.add_paragraph(); p2.text = subtitle
    p2.font.size = Pt(18); p2.font.color.rgb = ICE; p2.space_before = Pt(20)

def add_assessment(test_num, test_name, section, config, assessment, result, rcolor):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    tb = sl.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
    tb.fill.solid(); tb.fill.fore_color.rgb = NAVY; tb.line.fill.background()
    ttx = tb.text_frame; ttx.word_wrap = True; ttx.margin_left = Inches(0.6)
    p = ttx.paragraphs[0]; p.text = "Test %d: %s" % (test_num, test_name)
    p.font.size = Pt(28); p.font.bold = True; p.font.color.rgb = WHITE

    cx = sl.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(8.5), Inches(5.0))
    cf = cx.text_frame; cf.word_wrap = True
    labels = [("DWG Section:", section), ("Configuration:", config), ("Assessment:", "")]
    first = True
    for lbl, val in labels:
        if first:
            pa = cf.paragraphs[0]; first = False
        else:
            pa = cf.add_paragraph()
        pa.text = "%s %s" % (lbl, val)
        pa.font.size = Pt(14); pa.font.color.rgb = DKGRAY; pa.font.bold = True
        pa.space_before = Pt(8)
    pa = cf.add_paragraph(); pa.text = assessment
    pa.font.size = Pt(13); pa.font.color.rgb = DKGRAY; pa.space_before = Pt(4)

    rx = sl.shapes.add_shape(5, Inches(10.0), Inches(1.4), Inches(2.8), Inches(0.8))
    rx.fill.solid(); rx.fill.fore_color.rgb = rcolor; rx.line.fill.background()
    rtf = rx.text_frame; rtf.paragraphs[0].text = result
    rtf.paragraphs[0].font.size = Pt(22); rtf.paragraphs[0].font.bold = True
    rtf.paragraphs[0].font.color.rgb = WHITE; rtf.paragraphs[0].alignment = PP_ALIGN.CENTER
    rtf.vertical_anchor = MSO_ANCHOR.MIDDLE

def add_plot(test_num, test_name, plot_title, img1, img2=None):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    tb = sl.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.9))
    tb.fill.solid(); tb.fill.fore_color.rgb = NAVY; tb.line.fill.background()
    ttx = tb.text_frame; ttx.word_wrap = True; ttx.margin_left = Inches(0.6)
    p = ttx.paragraphs[0]; p.text = "Test %d: %s \u2014 %s" % (test_num, test_name, plot_title)
    p.font.size = Pt(22); p.font.bold = True; p.font.color.rgb = WHITE
    if img2 and os.path.exists(img2):
        if os.path.exists(img1):
            sl.shapes.add_picture(img1, Inches(0.3), Inches(1.1), Inches(6.3), Inches(5.8))
        sl.shapes.add_picture(img2, Inches(6.8), Inches(1.1), Inches(6.3), Inches(5.8))
    elif os.path.exists(img1):
        sl.shapes.add_picture(img1, Inches(1.5), Inches(1.0), Inches(10.3), Inches(6.0))

# ================================================================
now = datetime.datetime.now().strftime("%m/%d/%Y %I:%M %p")
add_title_slide("NLR GFM Inverter Model\nMQT Batch Results \u2014 Anti-Windup ON",
                "Tests 6, 7 | Vsched = 1.03 | Pset = 0.4 pu | " + now)

# Test 6: LVRT Dips IEEE 2800
d6 = os.path.join(PLOT_DIR, "test06_LVRT_Dips_IEEE_2800_AW")
add_assessment(6, "LVRT Dips IEEE 2800", "3.1.5.4",
    "Vsched = 1.03 | V-Reg ON | Anti-Windup ON | Pset = 0.4 pu",
    "Series of IEEE 2800 voltage dips (0.7, 0.0, 0.25, 0.5, 0.7 pu). Model rides through "
    "all dips with large transient P/Q excursions at each voltage transition (Ppoi swinging "
    "-1.1 to +1.0 pu momentarily). Reactive current injection observed during each dip. "
    "Power recovers to full output (Ppoi = 0.40 pu) after all events. Frequency dips "
    "to 53.2 Hz during the deepest voltage dips (brief transient). Final recovery is clean "
    "with Fpoi = 60.00 Hz. Current limiter not reported in 22-channel output. Transient "
    "magnitude larger than no-anti-windup case but model maintains ride-through capability.",
    "CONDITIONAL PASS", COND_AMBER)
add_plot(6, "LVRT Dips IEEE 2800", "Source Voltage & POI Power",
         os.path.join(d6, "vsource.png"), os.path.join(d6, "poi_power.png"))
add_plot(6, "LVRT Dips IEEE 2800", "Power-Voltage Overlay",
         os.path.join(d6, "pq_vrms_overlay.png"))
add_plot(6, "LVRT Dips IEEE 2800", "POI Frequency",
         os.path.join(d6, "fpoi.png"))
add_plot(6, "LVRT Dips IEEE 2800", "POI Current",
         os.path.join(d6, "ipoi.png"))
add_plot(6, "LVRT Dips IEEE 2800", "3-Phase Voltage",
         os.path.join(d6, "vinst_3phase.png"))
add_plot(6, "LVRT Dips IEEE 2800", "Model/PMU Power",
         os.path.join(d6, "model_pmu_power.png"))
add_plot(6, "LVRT Dips IEEE 2800", "Breaker Status",
         os.path.join(d6, "breaker_status.png"))

# Test 7: HVRT Preferred IEEE 2800
d7 = os.path.join(PLOT_DIR, "test07_HVRT_Preferred_IEEE_2800_AW")
add_assessment(7, "HVRT Preferred IEEE 2800", "3.1.5.5",
    "Vsched = 1.03 | V-Reg ON | Anti-Windup ON | Pset = 0.4 pu",
    "High-voltage steps to 1.2 pu at t~5s, then 1.1, 1.05 pu segments, extreme 1.6 pu "
    "at t~24s. With anti-windup current limiting, the model exhibits sustained P/Q "
    "oscillations during the entire HVRT ride-through period (t=5-24s), with P swinging "
    "-1.0 to +0.5 pu and Q oscillating similarly. Current limiter active ~14.4s. "
    "At t=24s (1.6 pu event), Qpoi spikes to ~3.3 pu. After voltage returns, power "
    "recovers to Ppoi ~0.41 pu. Same limiter-cycling oscillation pattern as Test 3 "
    "(HVRT Legacy). Degradation vs. no-anti-windup case (which was CONDITIONAL PASS "
    "with brief transient spikes but no sustained oscillations).",
    "FAIL", FAIL_RED)
add_plot(7, "HVRT Preferred IEEE 2800", "Source Voltage & POI Power",
         os.path.join(d7, "vsource.png"), os.path.join(d7, "poi_power.png"))
add_plot(7, "HVRT Preferred IEEE 2800", "Power-Voltage Overlay",
         os.path.join(d7, "pq_vrms_overlay.png"))
add_plot(7, "HVRT Preferred IEEE 2800", "POI Frequency",
         os.path.join(d7, "fpoi.png"))
add_plot(7, "HVRT Preferred IEEE 2800", "POI Current",
         os.path.join(d7, "ipoi.png"))
add_plot(7, "HVRT Preferred IEEE 2800", "3-Phase Voltage",
         os.path.join(d7, "vinst_3phase.png"))
add_plot(7, "HVRT Preferred IEEE 2800", "Model/PMU Power",
         os.path.join(d7, "model_pmu_power.png"))
add_plot(7, "HVRT Preferred IEEE 2800", "Breaker Status",
         os.path.join(d7, "breaker_status.png"))

ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
fname = "GFM_MQT_AW_Batch_Tests6_7_%s.pptx" % ts
prs.save(fname)
print("Saved: %s (%d slides)" % (fname, len(prs.slides)))
