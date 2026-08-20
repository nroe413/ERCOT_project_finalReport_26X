from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os, datetime

PLOT_DIR = "plots"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(0x1E, 0x27, 0x61)
ICE  = RGBColor(0xCA, 0xDC, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DKGRAY = RGBColor(0x33, 0x33, 0x33)
FAIL_RED = RGBColor(0xB8, 0x50, 0x42)

def add_title_slide(title, subtitle):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = NAVY
    tx = sl.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(2.5))
    tf = tx.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title
    p.font.size = Pt(40); p.font.bold = True; p.font.color.rgb = WHITE
    p2 = tf.add_paragraph(); p2.text = subtitle
    p2.font.size = Pt(18); p2.font.color.rgb = ICE; p2.space_before = Pt(20)

def add_assessment(test_num, test_name, section, config, assessment, result, rcolor):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    tb = sl.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
    tb.fill.solid(); tb.fill.fore_color.rgb = NAVY; tb.line.fill.background()
    ttx = tb.text_frame; ttx.word_wrap = True; ttx.margin_left = Inches(0.6)
    p = ttx.paragraphs[0]; p.text = "Test %d: %s" % (test_num, test_name)
    p.font.size = Pt(28); p.font.bold = True; p.font.color.rgb = WHITE

    cx = sl.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(8.5), Inches(5.0))
    cf = cx.text_frame; cf.word_wrap = True
    labels = [("DWG Section:", section), ("Configuration:", config), ("Assessment:", "")]
    first = True
    for lbl, val in labels:
        if first:
            pa = cf.paragraphs[0]; first = False
        else:
            pa = cf.add_paragraph()
        pa.text = "%s %s" % (lbl, val)
        pa.font.size = Pt(14); pa.font.color.rgb = DKGRAY; pa.font.bold = True
        pa.space_before = Pt(8)
    pa = cf.add_paragraph(); pa.text = assessment
    pa.font.size = Pt(13); pa.font.color.rgb = DKGRAY; pa.space_before = Pt(4)

    rx = sl.shapes.add_shape(5, Inches(10.0), Inches(1.4), Inches(2.8), Inches(0.8))
    rx.fill.solid(); rx.fill.fore_color.rgb = rcolor; rx.line.fill.background()
    rtf = rx.text_frame; rtf.paragraphs[0].text = result
    rtf.paragraphs[0].font.size = Pt(22); rtf.paragraphs[0].font.bold = True
    rtf.paragraphs[0].font.color.rgb = WHITE; rtf.paragraphs[0].alignment = PP_ALIGN.CENTER
    rtf.vertical_anchor = MSO_ANCHOR.MIDDLE

def add_plot(test_num, test_name, plot_title, img1, img2=None):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    tb = sl.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.9))
    tb.fill.solid(); tb.fill.fore_color.rgb = NAVY; tb.line.fill.background()
    ttx = tb.text_frame; ttx.word_wrap = True; ttx.margin_left = Inches(0.6)
    p = ttx.paragraphs[0]; p.text = "Test %d: %s \u2014 %s" % (test_num, test_name, plot_title)
    p.font.size = Pt(22); p.font.bold = True; p.font.color.rgb = WHITE
    if img2 and os.path.exists(img2):
        if os.path.exists(img1):
            sl.shapes.add_picture(img1, Inches(0.3), Inches(1.1), Inches(6.3), Inches(5.8))
        sl.shapes.add_picture(img2, Inches(6.8), Inches(1.1), Inches(6.3), Inches(5.8))
    elif os.path.exists(img1):
        sl.shapes.add_picture(img1, Inches(1.5), Inches(1.0), Inches(10.3), Inches(6.0))

# ================================================================
now = datetime.datetime.now().strftime("%m/%d/%Y %I:%M %p")
add_title_slide("NLR GFM Inverter Model\nMQT Batch Results \u2014 Anti-Windup ON",
                "Tests 3, 8, 9 | Pset = 0.4 pu | " + now)

# Test 3: HVRT ERCOT Legacy
d3 = os.path.join(PLOT_DIR, "test03_HVRT_ERCOT_Legacy_AW")
add_assessment(3, "HVRT ERCOT Legacy", "3.1.5.5",
    "Vsched = 1.03 | V-Reg ON | Anti-Windup ON | Pset = 0.4 pu",
    "High-voltage steps at t~5s through 1.2, 1.1, 1.05 pu segments, then extreme 2.0 pu "
    "at t~23s. With anti-windup current limiting enabled, the model exhibits sustained P/Q "
    "oscillations (Ppoi swinging -1.5 to +0.5 pu) during the entire HVRT ride-through period "
    "(t=5-23s). Current limiter active ~14.4s. After the 2.0 pu extreme event ends, power "
    "recovers to Ppoi ~0.4 pu. Sustained oscillations during ride-through represent a "
    "degradation vs. the no-anti-windup case (which showed brief transient spikes but no "
    "sustained oscillations).",
    "FAIL", FAIL_RED)
add_plot(3, "HVRT ERCOT Legacy", "Source Voltage & POI Power",
         os.path.join(d3, "vsource.png"), os.path.join(d3, "poi_power.png"))
add_plot(3, "HVRT ERCOT Legacy", "Power-Voltage Overlay",
         os.path.join(d3, "pq_vrms_overlay.png"))
add_plot(3, "HVRT ERCOT Legacy", "POI Frequency",
         os.path.join(d3, "fpoi.png"))
add_plot(3, "HVRT ERCOT Legacy", "POI Current",
         os.path.join(d3, "ipoi.png"))
add_plot(3, "HVRT ERCOT Legacy", "3-Phase Voltage",
         os.path.join(d3, "vinst_3phase.png"))
add_plot(3, "HVRT ERCOT Legacy", "Model/PMU Power",
         os.path.join(d3, "model_pmu_power.png"))
add_plot(3, "HVRT ERCOT Legacy", "Breaker Status",
         os.path.join(d3, "breaker_status.png"))

# Test 8: Phase Angle Jump (Down)
d8 = os.path.join(PLOT_DIR, "test08_Phase_Angle_Jump_Down_AW")
add_assessment(8, "Phase Angle Jump (Down)", "3.1.5.9",
    "Vsched = 1.01 | V-Reg ON | Anti-Windup ON | Pset = 0.4 pu",
    "Phase angle jumps at t=5, 10, 15, 20, 25s cause severe transients. After the second "
    "jump (t~10s), the model enters sustained P/Q oscillations (Ppoi swinging -1.0 to +1.0 pu) "
    "that persist for the remainder of the simulation. Frequency deviates sharply at each jump "
    "(min 49.4 Hz, max 68.0 Hz). Ppoi settles to -0.58 pu at end (reversed from Pset = 0.4). "
    "Similar failure mode to no-anti-windup case.",
    "FAIL", FAIL_RED)
add_plot(8, "Phase Angle Jump (Down)", "Source Voltage & POI Power",
         os.path.join(d8, "vsource.png"), os.path.join(d8, "poi_power.png"))
add_plot(8, "Phase Angle Jump (Down)", "Power-Voltage Overlay",
         os.path.join(d8, "pq_vrms_overlay.png"))
add_plot(8, "Phase Angle Jump (Down)", "POI Frequency",
         os.path.join(d8, "fpoi.png"))
add_plot(8, "Phase Angle Jump (Down)", "POI Current",
         os.path.join(d8, "ipoi.png"))
add_plot(8, "Phase Angle Jump (Down)", "3-Phase Voltage",
         os.path.join(d8, "vinst_3phase.png"))
add_plot(8, "Phase Angle Jump (Down)", "Model/PMU Power",
         os.path.join(d8, "model_pmu_power.png"))
add_plot(8, "Phase Angle Jump (Down)", "Breaker Status",
         os.path.join(d8, "breaker_status.png"))

# Test 9: Phase Angle Jump (Up)
d9 = os.path.join(PLOT_DIR, "test09_Phase_Angle_Jump_Up_AW")
add_assessment(9, "Phase Angle Jump (Up)", "3.1.5.9",
    "Vsched = 1.01 | V-Reg ON | Anti-Windup ON | Pset = 0.4 pu",
    "Same failure pattern as Test 8 (mirror image). Phase angle jumps cause severe P/Q "
    "transients with growing magnitude. Model enters sustained oscillations after successive "
    "jumps. Frequency spikes to 76.7 Hz. Current limiter active ~7.2s during transient window. "
    "Ppoi settles to -0.57 pu (reversed). Anti-windup current limiting does not improve phase "
    "angle jump response vs. baseline.",
    "FAIL", FAIL_RED)
add_plot(9, "Phase Angle Jump (Up)", "Source Voltage & POI Power",
         os.path.join(d9, "vsource.png"), os.path.join(d9, "poi_power.png"))
add_plot(9, "Phase Angle Jump (Up)", "Power-Voltage Overlay",
         os.path.join(d9, "pq_vrms_overlay.png"))
add_plot(9, "Phase Angle Jump (Up)", "POI Frequency",
         os.path.join(d9, "fpoi.png"))
add_plot(9, "Phase Angle Jump (Up)", "POI Current",
         os.path.join(d9, "ipoi.png"))
add_plot(9, "Phase Angle Jump (Up)", "3-Phase Voltage",
         os.path.join(d9, "vinst_3phase.png"))
add_plot(9, "Phase Angle Jump (Up)", "Model/PMU Power",
         os.path.join(d9, "model_pmu_power.png"))
add_plot(9, "Phase Angle Jump (Up)", "Breaker Status",
         os.path.join(d9, "breaker_status.png"))

ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
fname = "GFM_MQT_AW_Batch_Tests3_8_9_%s.pptx" % ts
prs.save(fname)
print("Saved: %s (%d slides)" % (fname, len(prs.slides)))
