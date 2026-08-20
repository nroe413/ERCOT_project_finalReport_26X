"""Build the Case 1 (Voltage Step Down) PowerPoint deck for the NLR model.

Mirrors the PNNL deck builder format; only the subtitle and output
filename differ so the eventual side-by-side comparison reads cleanly.
Inherits theme from slide_templates/meeting_update_042226_r3.pptx.
"""
import datetime
import glob
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
CASE_DIR = os.path.dirname(HERE)
PLOT_DIR = os.path.join(HERE, "plots")
TEMPLATE = os.path.join(
    CASE_DIR, "..", "slide_templates", "meeting_update_042226_r3.pptx"
)
TEMPLATE = os.path.abspath(TEMPLATE)

BODY_FONT = "Tw Cen MT"
DKGRAY = RGBColor(0x40, 0x40, 0x40)

SLIDE_W_IN = 10.0
SLIDE_H_IN = 5.625

CASE_TITLE = "Case 1: Voltage Step Down from 1.0 p.u. to 0.95 p.u."
CASE_SUBTITLE = "NLR Benchmark — PSCAD"
OUT_FNAME_PREFIX = "Case1_VoltageStepDown_NLR"

PLOTS = [
    ("V (pu)",    "v_pu"),
    ("I (pu)",    "i_pu"),
    ("P (pu)",    "p_pu"),
    ("Q (pu)",    "q_pu"),
    ("Freq (Hz)", "freq_hz"),
]


def latest_run_ts():
    marker = os.path.join(PLOT_DIR, "latest_run.txt")
    if os.path.exists(marker):
        with open(marker) as f:
            ts = f.read().strip()
        if ts:
            return ts
    cands = glob.glob(os.path.join(PLOT_DIR, "v_pu_*.png"))
    if not cands:
        return None
    cands.sort(key=os.path.getmtime, reverse=True)
    base = os.path.basename(cands[0])
    return base[len("v_pu_"):-len(".png")]


def load_template():
    prs = Presentation(TEMPLATE)
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        )
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)
    return prs


def layout_by_name(prs, name):
    for m in prs.slide_masters:
        for lyt in m.slide_layouts:
            if lyt.name == name:
                return lyt
    raise KeyError(name)


def first_layout(prs, *names):
    for n in names:
        try:
            return layout_by_name(prs, n)
        except KeyError:
            continue
    return prs.slide_layouts[0]


def force_white_background(slide):
    from lxml import etree
    from pptx.oxml.ns import qn

    cSld = slide.element.find(qn("p:cSld"))
    existing = cSld.find(qn("p:bg"))
    if existing is not None:
        cSld.remove(existing)
    bg = etree.SubElement(cSld, qn("p:bg"))
    cSld.remove(bg)
    cSld.insert(0, bg)
    bgPr = etree.SubElement(bg, qn("p:bgPr"))
    solidFill = etree.SubElement(bgPr, qn("a:solidFill"))
    srgb = etree.SubElement(solidFill, qn("a:srgbClr"))
    srgb.set("val", "FFFFFF")
    etree.SubElement(bgPr, qn("a:effectLst"))


def clear_all_placeholders(slide):
    to_remove = list(slide.placeholders)
    for shp in to_remove:
        sp = shp._element
        sp.getparent().remove(sp)


def add_custom_title(slide, text, *, top_in=0.25, height_in=0.70, size=18,
                     left_pad=0.4):
    width_in = SLIDE_W_IN - 2 * left_pad
    tb = slide.shapes.add_textbox(
        Inches(left_pad), Inches(top_in), Inches(width_in), Inches(height_in)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = BODY_FONT
    return tb


def image_size_at_width(path, width_in):
    im = Image.open(path)
    w_px, h_px = im.size
    return width_in * (h_px / w_px)


def add_picture_aspect(slide, path, left_in, top_in, width_in):
    return slide.shapes.add_picture(
        path, Inches(left_in), Inches(top_in), width=Inches(width_in)
    )


def add_title_slide(prs, title, subtitle):
    layout = first_layout(prs, "Title Slide", "Title and Content")
    sl = prs.slides.add_slide(layout)
    force_white_background(sl)
    clear_all_placeholders(sl)

    title_top = SLIDE_H_IN * 0.30
    tb = sl.shapes.add_textbox(
        Inches(0.5), Inches(title_top), Inches(SLIDE_W_IN - 1.0), Inches(1.2)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.name = BODY_FONT

    sb = sl.shapes.add_textbox(
        Inches(0.5), Inches(title_top + 1.3),
        Inches(SLIDE_W_IN - 1.0), Inches(0.7)
    )
    sbf = sb.text_frame
    sbf.word_wrap = True
    p2 = sbf.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = subtitle
    r2.font.size = Pt(18)
    r2.font.name = BODY_FONT
    r2.font.color.rgb = DKGRAY
    return sl


def add_plot_slide(prs, plot_title, img_path):
    layout = first_layout(prs, "Title and Content", "Title Only")
    sl = prs.slides.add_slide(layout)
    force_white_background(sl)
    clear_all_placeholders(sl)
    add_custom_title(sl, f"{CASE_TITLE} — {plot_title}", top_in=0.25,
                     height_in=0.70, size=18)

    content_top = 1.10
    content_h_max = SLIDE_H_IN - content_top - 0.30

    max_w = SLIDE_W_IN - 1.0
    h_at_max_w = image_size_at_width(img_path, max_w)
    if h_at_max_w <= content_h_max:
        w = max_w
        h = h_at_max_w
    else:
        im = Image.open(img_path)
        wpx, hpx = im.size
        h = content_h_max
        w = h * (wpx / hpx)

    left = (SLIDE_W_IN - w) / 2
    top = content_top + (content_h_max - h) / 2
    add_picture_aspect(sl, img_path, left, top, w)
    return sl


def main():
    prs = load_template()

    global SLIDE_W_IN, SLIDE_H_IN
    SLIDE_W_IN = prs.slide_width / 914400.0
    SLIDE_H_IN = prs.slide_height / 914400.0

    add_title_slide(prs, CASE_TITLE, CASE_SUBTITLE)

    plot_ts = latest_run_ts()
    if plot_ts is None:
        raise SystemExit("No plot PNGs found in plots/. Run _plot_vstepdown.py first.")
    print(f"Using plot run: {plot_ts}")

    for plot_title, base in PLOTS:
        img_path = os.path.join(PLOT_DIR, f"{base}_{plot_ts}.png")
        if not os.path.exists(img_path):
            print(f"  WARN missing {img_path} — skipping")
            continue
        add_plot_slide(prs, plot_title, img_path)

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    out = os.path.join(HERE, f"{OUT_FNAME_PREFIX}_{ts}.pptx")
    prs.save(out)
    print(f"Saved: {out} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
